import os, tempfile, base64
os.environ["DB_PATH"] = os.path.join(tempfile.mkdtemp(), "t.db")
os.environ["ADMIN_PASSWORD"] = "secret"
from io import BytesIO
from openpyxl import load_workbook
import app as A

c = A.app.test_client()
H = {"Authorization": "Basic " + base64.b64encode(b"admin:secret").decode()}
fict = [  # DOMAINE, ANCIENNETE, AGE, RECOUVREMENT, PAR, PERFORMANCE, APPRECIATION
    (4, 3, 2, 4, 1, 1, "La qualité du recouvrement améliore la performance, avec des virgules; des « guillemets » et\nun retour à la ligne."),
    (2, 4, 3, 3, 2, 0, "Les retards de paiement peuvent réduire les résultats."),
    (1, 1, 1, 1, 4, 0, "Avis court."),
    (7, 5, 5, 2, 3, 1, "Le lien est direct."),
    (6, 2, 4, 4, 1, 1, "Une bonne gestion protège le portefeuille."),
]
keys = A.VARS
for i, r in enumerate(fict):
    p = dict(zip(keys, r), token=f"tok-fictif-{i:02d}-abcdefghij")
    assert c.post("/api/soumettre", json=p).status_code == 201

# double envoi : même jeton, aucune ligne en plus
p = dict(zip(keys, fict[0]), token="tok-fictif-00-abcdefghij")
assert c.post("/api/soumettre", json=p).json.get("doublon") is True

# rejets : code hors liste, valeur manquante, deux valeurs, 0 non confondu avec manquant, texte vide
ok = dict(zip(keys, fict[0]), token="tok-rejet-000000000000")
for bad in [{"DOMAINE": 8}, {"DOMAINE": None}, {"AGE": [1, 2]}, {"PERFORMANCE": 2}, {"PERFORMANCE": None},
            {"PAR": "abc"}, {"APPRECIATION": "   "}]:
    assert c.post("/api/soumettre", json={**ok, **bad}).status_code == 422, bad

# admin protégé
assert c.get("/admin").status_code == 401 and c.get("/admin/export.xlsx").status_code == 401
assert "5</b>" in c.get("/admin", headers=H).get_data(as_text=True)
assert c.get("/").status_code == 200

wb = load_workbook(BytesIO(c.get("/admin/export.xlsx", headers=H).data))
assert wb.sheetnames == ["DONNEES", "CODEBOOK", "README"], wb.sheetnames
d = list(wb["DONNEES"].values)
assert list(d[0]) == ["ID_REPONDANT", "DATE_SAISIE", "DOMAINE", "ANCIENNETE", "AGE", "RECOUVREMENT", "PAR", "PERFORMANCE", "APPRECIATION"]
assert len(d) == 6
for i, (row, src) in enumerate(zip(d[1:], fict), 1):
    assert row[0] == f"REP{i:03d}" and len(row[1]) == 16 and row[1][4] == "-" and row[1][10] == " "
    assert tuple(row[2:8]) == src[:6] and all(isinstance(v, int) for v in row[2:8]), row
    assert row[8] == src[6]
assert d[2][7] == 0  # "Non" reste 0
cb = list(wb["CODEBOOK"].values)
assert list(cb[0]) == ["VARIABLE", "LIBELLE", "CODE", "MODALITE"]
assert len(cb) == 1 + 7 + 5 + 5 + 4 + 4 + 2
assert {(r[0], r[2]): r[3] for r in cb[1:]}[("PERFORMANCE", 0)] == "Non"
assert ("PERFORMANCE", 1, "Oui") in {(r[0], r[2], r[3]) for r in cb[1:]}
assert len({(r[0], r[2]) for r in cb[1:]}) == len(cb) - 1  # aucun code en double
assert len(list(wb["README"].values)) > 10
os.makedirs("/mnt/user-data/outputs/padme-collecte", exist_ok=True)
wb.save("/mnt/user-data/outputs/padme-collecte/exemple_export_fictif.xlsx")
print("TOUS LES TESTS OK")
for r in d: print(r[:8], "|", str(r[8])[:30].replace("\n", " "))
