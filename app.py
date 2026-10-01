"""Collecte de données PADME Pobè : formulaire, validation, SQLite, export .xlsx codifié."""
import os, secrets, sqlite3
from datetime import datetime, timezone, timedelta
from io import BytesIO
from flask import Flask, request, jsonify, send_file, Response, render_template_string
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

DB = os.environ.get("DB_PATH", "reponses.db")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD") or secrets.token_urlsafe(8)
if not os.environ.get("ADMIN_PASSWORD"):
    print("ADMIN_PASSWORD non défini, mot de passe temporaire :", ADMIN_PASSWORD)
APPRECIATION_OBLIGATOIRE = True  # passer à False pour rendre la question 7 facultative
TZ = timezone(timedelta(hours=1))  # Bénin, UTC+1 sans changement d'heure

SECTIONS = ["Informations sur le répondant", "Gestion du crédit", "Performance financière", "Appréciation générale"]
P = [(1, "0 % à moins de 25 %"), (2, "25 % à moins de 50 %"), (3, "50 % à moins de 75 %"), (4, "75 % à 100 %")]
Q = [
    dict(v="DOMAINE", s=0, lib="Domaine d’intervention", txt="Dans quel domaine intervenez-vous principalement au sein de PADME ?",
         m=[(1, "Évaluation / analyse des crédits"), (2, "Octroi / décision"), (3, "Suivi des crédits"), (4, "Recouvrement"),
            (5, "Comptabilité / Finance"), (6, "Gestion du portefeuille"), (7, "Gestion / supervision")]),
    dict(v="ANCIENNETE", s=0, lib="Ancienneté dans la fonction actuelle", txt="Depuis combien de temps exercez-vous votre fonction actuelle ?",
         m=[(1, "Moins d’un an"), (2, "1 à 3 ans"), (3, "4 à 6 ans"), (4, "Plus de 6 ans"), (5, "Préfère ne pas répondre")]),
    dict(v="AGE", s=0, lib="Tranche d’âge", txt="Dans quelle tranche d’âge vous situez-vous ?",
         m=[(1, "Moins de 25 ans"), (2, "25 à 34 ans"), (3, "35 à 44 ans"), (4, "45 à 54 ans"), (5, "55 ans et plus")]),
    dict(v="RECOUVREMENT", s=1, lib="Taux de recouvrement",
         txt="Quelle proportion du montant total des échéances arrivées à terme est effectivement recouvrée ?", m=P),
    dict(v="PAR", s=1, lib="Portefeuille à risque (retard de paiement)",
         txt="Selon votre appréciation, quelle proportion du portefeuille de crédits présente actuellement un retard de paiement ?", m=P),
    dict(v="PERFORMANCE", s=2, lib="La gestion du crédit améliore la performance financière",
         txt="Selon votre appréciation, la qualité de la gestion du crédit contribue-t-elle à l’amélioration de la performance financière de PADME ?",
         m=[(0, "Non"), (1, "Oui")]),
]
for n, q in enumerate(Q, 1):
    q["n"] = n
OPEN = dict(v="APPRECIATION", s=3, n=7, lib="Appréciation du lien gestion du crédit / performance financière",
            txt="Quelle appréciation faites-vous du lien entre la gestion du crédit et la performance financière de PADME ?")
VARS = [q["v"] for q in Q] + [OPEN["v"]]
HEADERS = ["ID_REPONDANT", "DATE_SAISIE"] + VARS

app = Flask(__name__)


def db():
    c = sqlite3.connect(DB)
    cols = ",".join(f'{q["v"]} INTEGER NOT NULL CHECK({q["v"]} IN ({",".join(str(c) for c, _ in q["m"])}))' for q in Q)
    c.execute(f"""CREATE TABLE IF NOT EXISTS reponses(id INTEGER PRIMARY KEY AUTOINCREMENT,
        token TEXT NOT NULL UNIQUE, date_saisie TEXT NOT NULL, {cols}, APPRECIATION TEXT {'NOT NULL' if APPRECIATION_OBLIGATOIRE else ''})""")
    return c


def valider(d):
    """Retourne (valeurs, erreurs). Une variable = une seule valeur scalaire parmi les codes autorisés."""
    err, out = {}, {}
    for q in Q:
        x = d.get(q["v"])
        ok = isinstance(x, int) and not isinstance(x, bool) or (isinstance(x, str) and x.isdigit())
        if not ok or int(x) not in [c for c, _ in q["m"]]:
            err[q["v"]] = "Réponse obligatoire (une seule modalité)."
        else:
            out[q["v"]] = int(x)
    t = d.get("APPRECIATION")
    t = t.strip() if isinstance(t, str) else None
    if not t and APPRECIATION_OBLIGATOIRE:
        err["APPRECIATION"] = "Réponse obligatoire."
    out["APPRECIATION"] = t or None
    return out, err


@app.post("/api/soumettre")
def soumettre():
    d = request.get_json(silent=True) or {}
    tok = d.get("token")
    if not isinstance(tok, str) or not 16 <= len(tok) <= 64:
        return jsonify(ok=False, erreurs={"token": "Jeton invalide."}), 400
    vals, err = valider(d)
    if err:
        return jsonify(ok=False, erreurs=err), 422
    c = db()
    try:
        with c:
            c.execute(f"INSERT INTO reponses(token,date_saisie,{','.join(VARS)}) VALUES(?,?,{','.join('?' * len(VARS))})",
                      [tok, datetime.now(TZ).strftime("%Y-%m-%d %H:%M")] + [vals[v] for v in VARS])
    except sqlite3.IntegrityError:  # même jeton : double envoi, déjà enregistré
        return jsonify(ok=True, doublon=True), 200
    finally:
        c.close()
    return jsonify(ok=True), 201


def construire_xlsx():
    c = db()
    rows = c.execute(f"SELECT id,date_saisie,{','.join(VARS)} FROM reponses ORDER BY id").fetchall()
    c.close()
    wb, hf, fill = Workbook(), Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="12355B")

    def feuille(ws, head, data, widths):
        ws.append(head)
        for r in data:
            ws.append(r)
        for cell in ws[1]:
            cell.font, cell.fill = hf, fill
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "DONNEES"
    feuille(ws, HEADERS, [[f"REP{r[0]:03d}", r[1], *r[2:]] for r in rows], [16, 18] + [14] * 5 + [14, 60])
    feuille(wb.create_sheet("CODEBOOK"), ["VARIABLE", "LIBELLE", "CODE", "MODALITE"],
            [[q["v"], q["lib"], c, m] for q in Q for c, m in q["m"]], [16, 52, 8, 40])
    desc = [
        ["ID_REPONDANT", "Identifiant unique du répondant (REP001, REP002, ...), attribué dans l’ordre de soumission."],
        ["DATE_SAISIE", "Date et heure de soumission, format AAAA-MM-JJ HH:MM, heure du Bénin (UTC+1), enregistrée en texte."],
        *[[q["v"], f'{q["lib"]}. Variable catégorielle : codes dans la feuille CODEBOOK.'] for q in Q],
        ["APPRECIATION", "Question ouverte : texte intégral saisi par le répondant."],
        ["Codification", "Les colonnes de la feuille DONNEES contiennent uniquement les codes numériques, jamais les libellés."],
        ["Une seule modalité", "Chaque question à choix unique enregistre une seule valeur dans une seule colonne."],
        ["Valeurs manquantes", "Toutes les questions étant obligatoires, aucune cellule vide n’est attendue. Aucune absence n’est remplacée par 0 (0 = Non pour PERFORMANCE)."],
        ["Analyse statistique", "Une ligne par répondant, une colonne par variable, noms sans espaces : importable dans SPSS, Stata, R ou Python."],
    ]
    feuille(wb.create_sheet("README"), ["ELEMENT", "EXPLICATION"], desc, [24, 110])
    b = BytesIO()
    wb.save(b)
    b.seek(0)
    return b, len(rows)


def admin_ok():
    a = request.authorization
    return bool(a and secrets.compare_digest((a.password or "").encode(), ADMIN_PASSWORD.encode()))


def refus():
    return Response("Accès réservé.", 401, {"WWW-Authenticate": 'Basic realm="Admin"'})


@app.get("/admin")
def admin():
    if not admin_ok():
        return refus()
    c = db()
    n, last = c.execute("SELECT COUNT(*),MAX(date_saisie) FROM reponses").fetchone()
    c.close()
    return render_template_string("""<!doctype html><meta charset=utf-8><meta name=viewport content="width=device-width,initial-scale=1">
<body style="font:16px system-ui;max-width:480px;margin:40px auto;padding:16px"><h1 style="color:#12355b">Administration</h1>
<p style="font-size:2.4rem;margin:.2em 0"><b>{{n}}</b> réponse(s)</p><p>Dernière : {{last or '—'}}</p>
<a href="/admin/export.xlsx" style="display:inline-block;background:#12355b;color:#fff;padding:14px 22px;border-radius:8px;text-decoration:none">Télécharger (.xlsx)</a></body>""", n=n, last=last)


@app.get("/admin/export.xlsx")
def export():
    if not admin_ok():
        return refus()
    b, _ = construire_xlsx()
    nom = f"padme_donnees_{datetime.now(TZ).strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(b, as_attachment=True, download_name=nom,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


PAGE = """<!doctype html><html lang="fr"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Enquête PADME Pobè</title><style>
:root{--b:#12355b}*{box-sizing:border-box}body{margin:0;font:16px/1.5 system-ui,sans-serif;background:#f3f5f8;color:#1b2430}
main{max-width:640px;margin:0 auto;padding:16px}h1{font-size:1.25rem;color:var(--b);margin:.2rem 0}h2{font-size:1.05rem;color:var(--b)}.note{font-size:.85rem;color:#566;margin:0}
.bar{height:8px;background:#d8dee6;border-radius:4px;overflow:hidden;margin:10px 0 4px}.bar i{display:block;height:100%;width:0;background:var(--b);transition:width .2s}
.q{background:#fff;border:1px solid #d8dee6;border-radius:10px;padding:14px;margin:12px 0}.q.err{border-color:#b3261e}.q p{margin:0 0 8px;font-weight:600}
.m{display:none;color:#b3261e;font-size:.85rem}.err .m{display:block}
label.o{display:flex;gap:10px;align-items:center;min-height:48px;padding:6px 10px;border:1px solid #d8dee6;border-radius:8px;margin:6px 0;cursor:pointer}
label.o:has(input:checked){border-color:var(--b);background:#e8eef6}input[type=radio]{width:22px;height:22px;accent-color:var(--b);flex:none}
textarea{width:100%;min-height:130px;font:inherit;padding:10px;border:1px solid #c5ccd6;border-radius:8px}
button{font:inherit;font-weight:600;min-height:48px;padding:0 22px;border:0;border-radius:8px;background:var(--b);color:#fff}button.s{background:#fff;color:var(--b);border:1px solid var(--b)}button:disabled{opacity:.5}
.nav{display:flex;justify-content:space-between;gap:10px;margin-top:8px}[hidden]{display:none!important}
</style><main><h1>Enquête : gestion du crédit et performance financière</h1>
<p class="note">Outil de collecte académique (mémoire), PADME agence de Pobè. Réponses anonymes, environ 2 minutes.</p>
<div class="bar"><i id="bar"></i></div><p class="note" id="lbl"></p>
<div id="f">{% for si in range(4) %}<section class="sec"><h2>{{S[si]}}</h2>
{% for q in Q if q.s==si %}<div class="q" data-v="{{q.v}}" data-k="r"><p>{{q.n}}. {{q.txt}}</p>
{% for c,t in q.m %}<label class="o"><input type="radio" name="{{q.v}}" value="{{c}}"> {{t}}</label>{% endfor %}
<span class="m">Veuillez choisir une seule réponse.</span></div>{% endfor %}
{% if si==3 %}<div class="q" data-v="APPRECIATION" data-k="t"><p>{{O.n}}. {{O.txt}}</p><textarea></textarea><span class="m">Veuillez saisir votre réponse.</span></div>{% endif %}
</section>{% endfor %}
<p class="m" id="gerr" style="display:block;min-height:1.2em"></p>
<div class="nav"><button class="s" id="prev">Précédent</button><button id="next">Suivant</button><button id="sub" hidden>Soumettre</button></div></div>
<div id="ok" hidden><div class="q"><h2>Merci pour votre participation</h2><p style="font-weight:400">Votre réponse a bien été enregistrée. Vous pouvez fermer cette page.</p></div></div></main>
<script>
const S=[...document.querySelectorAll('.sec')],$=id=>document.getElementById(id),token=(crypto.randomUUID?crypto.randomUUID():Date.now()+'-'+Math.random().toString(36).slice(2)+'0000000000');
let i=0,sending=false;
function get(q){if(q.dataset.k=='r'){const r=q.querySelector('input:checked');return r?Number(r.value):null}const t=q.querySelector('textarea').value.trim();return t||null}
function show(n){i=n;S.forEach((s,k)=>s.hidden=k!=n);$('bar').style.width=((n+1)/S.length*100)+'%';$('lbl').textContent='Étape '+(n+1)+' sur '+S.length;
$('prev').hidden=n==0;$('next').hidden=n==S.length-1;$('sub').hidden=n!=S.length-1;scrollTo(0,0)}
function check(n){let good=true;S[n].querySelectorAll('.q').forEach(q=>{const bad=get(q)===null;q.classList.toggle('err',bad);if(bad)good=false});$('gerr').textContent=good?'':'Merci de répondre à toutes les questions avant de continuer.';return good}
$('next').onclick=()=>check(i)&&show(i+1);$('prev').onclick=()=>show(i-1);
$('sub').onclick=async()=>{if(sending||!check(i))return;sending=true;$('sub').disabled=true;
const body={token};document.querySelectorAll('.q').forEach(q=>body[q.dataset.v]=get(q));
try{const r=await fetch('/api/soumettre',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(body)});
if(r.ok){$('f').hidden=true;$('ok').hidden=false;$('bar').style.width='100%';$('lbl').textContent='Terminé';return}
$('gerr').textContent='Certaines réponses sont invalides, veuillez vérifier.'}catch(e){$('gerr').textContent='Connexion impossible, réessayez.'}
sending=false;$('sub').disabled=false};
show(0);
</script></html>"""


@app.get("/")
def index():
    return render_template_string(PAGE, Q=Q, O=OPEN, S=SECTIONS)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
